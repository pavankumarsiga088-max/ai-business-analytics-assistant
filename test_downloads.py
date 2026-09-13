"""
Unit Tests for Downloads Module (Phase 11C).

Validates:
1. DataFrame -> CSV bytes
2. CSV contains expected column names
3. CSV contains expected rows
4. CSV does not contain an unwanted index column
5. DataFrame -> Excel bytes
6. Excel file can be reopened with openpyxl
7. Excel contains expected columns
8. Excel contains expected rows
9. Plotly Figure -> HTML bytes
10. HTML contains Plotly content
11. Empty DataFrame handling (CSV and Excel)
12. DataFrame immutability after CSV export
13. DataFrame immutability after Excel export
14. Figure remains valid after HTML export
15. Filename sanitization (path traversal, reserved characters, collapsing)
16. Unicode/special characters preserved across exports
"""

import io
import unittest
import openpyxl
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go

from src.analytics.downloads import (
    dataframe_to_csv_bytes,
    dataframe_to_excel_bytes,
    figure_to_html_bytes,
    sanitize_filename,
    build_download_filename,
)


class TestDownloads(unittest.TestCase):

    def setUp(self):
        self.sample_df = pd.DataFrame({
            "region": ["East", "South", "North"],
            "total_revenue": [207141.50, 159251.00, 112362.75],
            "order_count": [45, 38, 26],
        })
        self.sample_fig = px.bar(
            self.sample_df,
            x="region",
            y="total_revenue",
            title="Revenue by Region",
        )

    # -----------------------------------------------------------------
    # 1. DataFrame -> CSV bytes
    # -----------------------------------------------------------------
    def test_1_dataframe_to_csv_bytes(self):
        """Converting a DataFrame to CSV returns non-empty bytes."""
        csv_bytes = dataframe_to_csv_bytes(self.sample_df)
        self.assertIsInstance(csv_bytes, bytes)
        self.assertGreater(len(csv_bytes), 0)

    # -----------------------------------------------------------------
    # 2. CSV contains expected column names
    # -----------------------------------------------------------------
    def test_2_csv_contains_expected_columns(self):
        """CSV output contains all original column headers."""
        csv_text = dataframe_to_csv_bytes(self.sample_df).decode("utf-8")
        header_line = csv_text.splitlines()[0]
        self.assertIn("region", header_line)
        self.assertIn("total_revenue", header_line)
        self.assertIn("order_count", header_line)

    # -----------------------------------------------------------------
    # 3. CSV contains expected rows
    # -----------------------------------------------------------------
    def test_3_csv_contains_expected_rows(self):
        """CSV output contains all rows and formatted data."""
        csv_text = dataframe_to_csv_bytes(self.sample_df).decode("utf-8")
        lines = [line.strip() for line in csv_text.splitlines() if line.strip()]
        # 1 header + 3 data rows
        self.assertEqual(len(lines), 4)
        self.assertTrue(any("East" in line and "207141.5" in line for line in lines))
        self.assertTrue(any("South" in line and "159251.0" in line for line in lines))
        self.assertTrue(any("North" in line and "112362.75" in line for line in lines))

    # -----------------------------------------------------------------
    # 4. CSV does not contain an unwanted index column
    # -----------------------------------------------------------------
    def test_4_csv_excludes_index_column(self):
        """CSV output must not prepend a pandas 0-indexed column."""
        csv_text = dataframe_to_csv_bytes(self.sample_df).decode("utf-8")
        header_line = csv_text.splitlines()[0]
        self.assertEqual(header_line, "region,total_revenue,order_count")

        reloaded_df = pd.read_csv(io.StringIO(csv_text))
        self.assertListEqual(list(reloaded_df.columns), ["region", "total_revenue", "order_count"])
        self.assertEqual(len(reloaded_df), 3)

    # -----------------------------------------------------------------
    # 5. DataFrame -> Excel bytes
    # -----------------------------------------------------------------
    def test_5_dataframe_to_excel_bytes(self):
        """Converting a DataFrame to Excel returns valid bytes."""
        excel_bytes = dataframe_to_excel_bytes(self.sample_df)
        self.assertIsInstance(excel_bytes, bytes)
        self.assertGreater(len(excel_bytes), 0)

    # -----------------------------------------------------------------
    # 6. Excel file can be reopened with openpyxl
    # -----------------------------------------------------------------
    def test_6_excel_file_reopens_with_openpyxl(self):
        """The in-memory Excel bytes must parse cleanly as a valid workbook."""
        excel_bytes = dataframe_to_excel_bytes(self.sample_df)
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertIn("Query Results", wb.sheetnames)

    # -----------------------------------------------------------------
    # 7. Excel contains expected columns
    # -----------------------------------------------------------------
    def test_7_excel_contains_expected_columns(self):
        """Excel sheet header matches DataFrame column names."""
        excel_bytes = dataframe_to_excel_bytes(self.sample_df)
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        sheet = wb.active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, ["region", "total_revenue", "order_count"])

    # -----------------------------------------------------------------
    # 8. Excel contains expected rows
    # -----------------------------------------------------------------
    def test_8_excel_contains_expected_rows(self):
        """Excel workbook contains all data rows without index column."""
        excel_bytes = dataframe_to_excel_bytes(self.sample_df)
        reloaded_df = pd.read_excel(io.BytesIO(excel_bytes))
        self.assertEqual(len(reloaded_df), 3)
        self.assertListEqual(list(reloaded_df.columns), ["region", "total_revenue", "order_count"])
        self.assertEqual(reloaded_df.iloc[0]["region"], "East")
        self.assertAlmostEqual(float(reloaded_df.iloc[0]["total_revenue"]), 207141.50)

    # -----------------------------------------------------------------
    # 9. Plotly Figure -> HTML bytes
    # -----------------------------------------------------------------
    def test_9_figure_to_html_bytes(self):
        """Exporting a Plotly Figure returns valid HTML bytes."""
        html_bytes = figure_to_html_bytes(self.sample_fig)
        self.assertIsInstance(html_bytes, bytes)
        self.assertGreater(len(html_bytes), 0)

    # -----------------------------------------------------------------
    # 10. HTML contains Plotly content
    # -----------------------------------------------------------------
    def test_10_html_contains_plotly_content(self):
        """HTML output is a complete document containing Plotly script and figure data."""
        html_text = figure_to_html_bytes(self.sample_fig).decode("utf-8")
        self.assertIn("<html", html_text.lower())
        self.assertIn("plotly", html_text.lower())
        self.assertIn("Revenue by Region", html_text)
        self.assertIn("East", html_text)

    # -----------------------------------------------------------------
    # 11. Empty DataFrame handling
    # -----------------------------------------------------------------
    def test_11_empty_dataframe_handling(self):
        """Handling of None and empty DataFrames without exceptions."""
        self.assertEqual(dataframe_to_csv_bytes(None), b"")
        self.assertEqual(dataframe_to_excel_bytes(None), b"")
        self.assertEqual(figure_to_html_bytes(None), b"")

        empty_df = pd.DataFrame()
        self.assertEqual(dataframe_to_csv_bytes(empty_df), b"")

        # Empty DataFrame with defined columns
        col_empty_df = pd.DataFrame(columns=["a", "b"])
        csv_empty = dataframe_to_csv_bytes(col_empty_df).decode("utf-8").strip()
        self.assertEqual(csv_empty, "a,b")

    # -----------------------------------------------------------------
    # 12. DataFrame immutability after CSV export
    # -----------------------------------------------------------------
    def test_12_dataframe_immutability_csv(self):
        """Input DataFrame is completely unaltered after CSV export."""
        df_copy = self.sample_df.copy(deep=True)
        _ = dataframe_to_csv_bytes(self.sample_df)
        pd.testing.assert_frame_equal(self.sample_df, df_copy)

    # -----------------------------------------------------------------
    # 13. DataFrame immutability after Excel export
    # -----------------------------------------------------------------
    def test_13_dataframe_immutability_excel(self):
        """Input DataFrame is completely unaltered after Excel export."""
        df_copy = self.sample_df.copy(deep=True)
        _ = dataframe_to_excel_bytes(self.sample_df)
        pd.testing.assert_frame_equal(self.sample_df, df_copy)

    # -----------------------------------------------------------------
    # 14. Figure remains valid after HTML export
    # -----------------------------------------------------------------
    def test_14_figure_immutability_html(self):
        """Plotly Figure remains valid and inspectable after HTML export."""
        _ = figure_to_html_bytes(self.sample_fig)
        self.assertIsInstance(self.sample_fig, go.Figure)
        self.assertEqual(self.sample_fig.layout.title.text, "Revenue by Region")

    # -----------------------------------------------------------------
    # 15. Filename sanitization
    # -----------------------------------------------------------------
    def test_15_filename_sanitization(self):
        """Sanitizes path traversal, illegal characters, and generates filenames."""
        self.assertEqual(sanitize_filename("../../etc/passwd"), "etc_passwd")
        self.assertEqual(sanitize_filename("sales:data*2024?.csv"), "sales_data_2024")
        self.assertEqual(sanitize_filename('<script>alert("xss")</script>'), "script_alert_xss_script")
        self.assertEqual(sanitize_filename("   "), "analytics_result")
        self.assertEqual(sanitize_filename(None), "analytics_result")

        # Filename builder
        csv_fn = build_download_filename("result", "csv")
        self.assertEqual(csv_fn, "result.csv")

        excel_fn = build_download_filename("analysis", "xlsx", prefix="sales_data")
        self.assertEqual(excel_fn, "sales_data_analysis.xlsx")

        html_fn = build_download_filename("chart", "html")
        self.assertEqual(html_fn, "chart.html")

    # -----------------------------------------------------------------
    # 16. Unicode/text values are preserved where practical
    # -----------------------------------------------------------------
    def test_16_unicode_preservation(self):
        """Unicode characters (e.g. currency symbols, accented names) are preserved."""
        unicode_df = pd.DataFrame({
            "product_name": ["Café Crème", "München Spezial", "東京 Widget"],
            "price_eur": [4.50, 12.00, 25.00],
        })
        csv_bytes = dataframe_to_csv_bytes(unicode_df)
        csv_text = csv_bytes.decode("utf-8")
        self.assertIn("Café Crème", csv_text)
        self.assertIn("München Spezial", csv_text)
        self.assertIn("東京 Widget", csv_text)

        excel_bytes = dataframe_to_excel_bytes(unicode_df)
        reloaded_excel = pd.read_excel(io.BytesIO(excel_bytes))
        self.assertEqual(reloaded_excel.iloc[0]["product_name"], "Café Crème")
        self.assertEqual(reloaded_excel.iloc[2]["product_name"], "東京 Widget")


if __name__ == "__main__":
    unittest.main()
